"""Assemble the IEEE docx: single-column title block, two-column body,
full-width spanning sections for wide tables and figures."""
import json, sys, copy, os, re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

body_docx, front_json, out = sys.argv[1], sys.argv[2], sys.argv[3]
override = os.path.join(os.path.dirname(front_json), 'title_override.txt')

fm = json.load(open(front_json))
if os.path.exists(override):
    fm['title'] = open(override).read().strip()

doc = Document(body_docx)
body = doc.element.body
EMU_IN = 914400
SPAN_MIN_COLS = 5          # tables this wide get the full page width
SPAN_MIN_EMU = int(3.6 * EMU_IN)

# ---------------------------------------------------------------- front matter
def para(style, runs):
    p = doc.add_paragraph(style=style)
    for text, b, i in runs:
        r = p.add_run(text)
        if b is not None: r.bold = b
        if i is not None: r.italic = i
    return p

front = [
    para('IEEE Title',    [(fm['title'], None, None)]),
    para('IEEE Author',   [(fm['author'], None, None)]),
    para('IEEE Affil',    [(fm['affil'], None, None)]),
    para('IEEE Abstract', [('Abstract—', True, True), (fm['abstract'], True, False)]),
    para('IEEE Terms',    [('Index Terms—', True, True), (fm['keywords'], False, True)]),
]
for i, p in enumerate(front):
    body.remove(p._p)
    body.insert(i, p._p)

# ------------------------------------------------------------------- sections
final_sectPr = body.find(qn('w:sectPr'))
# Without an explicit type the final section defaults to nextPage, which would
# push the two-column body off the title page.
ft = final_sectPr.find(qn('w:type'))
if ft is None:
    ft = OxmlElement('w:type'); final_sectPr.insert(0, ft)
ft.set(qn('w:val'), 'continuous')

def sect(ncols):
    sp = copy.deepcopy(final_sectPr)
    for tag in ('w:headerReference', 'w:footerReference'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    cols = sp.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); sp.append(cols)
    cols.set(qn('w:num'), str(ncols))
    if ncols == 1:
        cols.attrib.pop(qn('w:space'), None)
    else:
        cols.set(qn('w:space'), str(int(0.25 * 1440)))
    return sp

# Title block is section 1: single column, breaking continuously into the body.
front[-1]._p.get_or_add_pPr().append(sect(1))

def spacer(sectPr):
    """A near-invisible paragraph that carries a section break."""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    for tag, val in (('w:spacing', None), ('w:rPr', None)):
        pass
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0'); sp.set(qn('w:line'), '20')
    sp.set(qn('w:lineRule'), 'exact')
    pPr.append(sp)
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '2'); rPr.append(sz)
    pPr.append(rPr)
    pPr.append(sectPr)
    p.append(pPr)
    return p

def span(first_el, last_el):
    """Put [first_el .. last_el] into a full-width single-column section."""
    children = list(body)
    i, j = children.index(first_el), children.index(last_el)
    body.insert(j + 1, spacer(sect(1)))   # closes the 1-col span
    body.insert(i, spacer(sect(2)))       # closes the preceding 2-col run

# -------------------------------------------------- find wide tables & figures
def prev_para(el):
    prev = el.getprevious()
    while prev is not None and not prev.tag.endswith('}p'):
        prev = prev.getprevious()
    return prev

def next_para(el):
    nxt = el.getnext()
    while nxt is not None and not nxt.tag.endswith('}p'):
        nxt = nxt.getnext()
    return nxt

def text_of(p):
    return ''.join(t.text or '' for t in p.iter(qn('w:t')))

targets = []
for tbl in body.findall(qn('w:tbl')):
    ncols = len(tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol')))
    if ncols >= SPAN_MIN_COLS:
        cap = prev_para(tbl)
        first = cap if cap is not None and re.match(r'\s*Table \d', text_of(cap)) else tbl
        targets.append((first, tbl))

for p in body.findall(qn('w:p')):
    ext = p.find('.//' + qn('wp:extent'))
    if ext is not None and int(ext.get('cx')) > SPAN_MIN_EMU:
        cap = next_para(p)
        last = cap if cap is not None and re.match(r'\s*Figure \d', text_of(cap)) else p
        targets.append((p, last))

# Apply back-to-front so earlier positions stay valid.
targets.sort(key=lambda t: list(body).index(t[0]), reverse=True)
for first, last in targets:
    span(first, last)

# ---------------------------------------- let Word size columns to content
for t in doc.tables:
    tblPr = t._tbl.tblPr
    for tag in ('w:tblLayout', 'w:tblW'):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'autofit')
    tblPr.append(layout)
    w = OxmlElement('w:tblW')
    w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), '5000')   # 100% of the section
    tblPr.append(w)
    for row in t.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for el in tcPr.findall(qn('w:tcW')):
                tcPr.remove(el)

# ------------------------------------------------------ IEEE 8pt table bodies
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = None
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.name = 'Times New Roman'

doc.save(out)
print(f'assembled -> {out}  ({len(targets)} full-width spans)')
